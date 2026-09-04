from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document


def _clean(text: str) -> str:
    text = re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()
    return text


def _read_docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    # 1. Header text FIRST (paragraphs, tables, and text boxes)
    for section in doc.sections:
        for container in [
            section.header,
            section.first_page_header,
            section.even_page_header,
        ]:
            if container:
                for paragraph in container.paragraphs:
                    if paragraph.text.strip():
                        parts.append(paragraph.text.strip())
                for table in container.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                parts.append(cell.text.strip())
                if container._element is not None:
                    for txbx in container._element.xpath(".//*[local-name()='txbxContent']"):
                        for p in txbx.xpath(".//*[local-name()='p']"):
                            txt = "".join(node.text for node in p.xpath(".//*[local-name()='t']") if node.text)
                            if txt.strip():
                                parts.append(txt.strip())

    # 2. Body paragraphs SECOND
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    # 3. Body tables THIRD
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())

    # 4. Body Text Boxes FOURTH
    if doc._element is not None:
        for txbx in doc._element.xpath(".//*[local-name()='txbxContent']"):
            for p in txbx.xpath(".//*[local-name()='p']"):
                txt = "".join(node.text for node in p.xpath(".//*[local-name()='t']") if node.text)
                if txt.strip():
                    parts.append(txt.strip())

    # 5. Footer text FIFTH
    for section in doc.sections:
        for container in [
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]:
            if container:
                for paragraph in container.paragraphs:
                    if paragraph.text.strip():
                        parts.append(paragraph.text.strip())
                for table in container.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                parts.append(cell.text.strip())

    return _clean("\n".join(parts))


def _after_label(text: str, label: str, stop_labels: list[str]) -> str:
    stops = "|".join(re.escape(item) for item in stop_labels)
    pattern = rf"{re.escape(label)}\s*:?\s*(.*?)(?={stops}|$)"
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return _clean(match.group(1)) if match else ""


def _first(pattern: str, text: str) -> str:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return _clean(match.group(1)) if match else ""


def _strip_status(value: str) -> str:
    value = re.sub(r"\s+Phù hợp\s+Không phù hợp.*$", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\s+Không phù hợp.*$", "", value, flags=re.IGNORECASE)
    return _clean(value)


def _clean_deposit_confirmation(value: str) -> str:
    value = re.sub(r"\bphế liệu nhập khẩu số\b\s*:?", " ", value, flags=re.IGNORECASE)
    value = re.sub(r"^số\s*:?\s*", "", value, flags=re.IGNORECASE)
    return _clean(value)


@dataclass(frozen=True)
class CertificateData:
    certificate_no: str = ""
    certificate_date: str = ""
    customer_name: str = ""
    customer_address: str = ""
    inspection_location: str = ""
    inspection_date: str = ""
    contract_no: str = ""
    invoice_no: str = ""
    bill_of_lading_no: str = ""
    deposit_confirmation: str = ""
    customs_declaration_no: str = ""
    product_name: str = ""
    hs_code: str = ""
    quantity: str = ""
    container_count: str = ""
    container_count_short: str = ""

    def as_dict(self) -> dict[str, str]:
        return self.__dict__.copy()


def _extract_cert_no(text: str) -> str:
    no = _first(r"Số\s+(?:chứng\s+thư|chứng\s+nhận|chứng\s+từ|CT)\s*:?\s*([A-Za-z0-9/.\-_xX]+)", text)
    if no:
        return no
    m_top = re.search(r"^.*?Số\s*:?\s*([A-Za-z0-9/.\-_xX]+)", text, re.IGNORECASE)
    if m_top and m_top.group(1).strip() and not m_top.group(1).strip().lower().startswith(("lượng", "container")):
        return _clean(m_top.group(1))
    m2 = re.search(r"Số\s*:?\s*([A-Za-z0-9/.\-_xX]+)(?=\s+.*?(?:Hải|Hà|TP|ngày|Tên|1\.|Địa|Hợp))", text, re.IGNORECASE)
    if m2 and m2.group(1).strip() and not m2.group(1).strip().lower().startswith(("lượng", "container")):
        return _clean(m2.group(1))
    m3 = re.search(r"Số\s*:?\s*([A-Za-z0-9/.\-_xX]{5,30})", text, re.IGNORECASE)
    if m3 and m3.group(1).strip():
        return _clean(m3.group(1))
    return ""


def _extract_cert_date(text: str) -> str:
    m1 = re.search(r"(?:Hải\s+Phòng|Hà\s+Nội|TP\.?\s*HCM|TP\.?\s*Hồ\s+Chí\s+Minh|ngày)\s*,\s*(ngày\s+[\d\w]{1,4}\s+tháng\s+[\d\w]{1,4}\s+năm\s+[\d\w]{4})", text, re.IGNORECASE)
    if m1 and m1.group(1).strip():
        return _clean(m1.group(1))
    m2 = re.search(r"(ngày\s+[\d\w]{1,4}\s+tháng\s+[\d\w]{1,4}\s+năm\s+[\d\w]{4})", text, re.IGNORECASE)
    if m2 and m2.group(1).strip():
        return _clean(m2.group(1))
    m3 = re.search(r"Ngày\s+(?:cấp\s+chứng\s+thư|cấp|chứng\s+thư|chứng\s+từ)\s*:?\s*([0-9/ ]{8,20})", text, re.IGNORECASE)
    if m3 and m3.group(1).strip():
        return _clean(m3.group(1))
    m4 = re.search(r"(\d{1,2}/\d{1,2}/\d{4})", text, re.IGNORECASE)
    if m4 and m4.group(1).strip():
        return _clean(m4.group(1))
    return ""


def extract_source_documents(paths: list[Path]) -> CertificateData:
    text = _clean(" ".join(_read_docx_text(path) for path in paths))
    customer_name = (
        _first(r"Tên khách hàng\s*:?\s*(.*?)(?=NỘI DUNG|NỘI\s+DUNG|1\.)", text)
        or _first(r"-\s*Tên tổ chức\s*:?\s*(.*?)(?=-\s*Địa chỉ|- Giấy|Giấy)", text)
    )
    customer_address = _first(r"-\s*Địa chỉ\s*:?\s*(.*?)(?=-\s*Giấy|Giấy xác nhận)", text)
    certificate_no = _extract_cert_no(text)
    certificate_date = _extract_cert_date(text)
    inspection_location = (
        _first(r"Địa điểm giám định/lấy mẫu\s*:?\s*(.*?)(?=7\.|Số lượng|8\.)", text)
        or _first(r"tại\s+([^,]+),\s*chúng tôi tiến hành", text)
    )
    contract_no = _strip_status(_first(r"Hợp đồng số\s*:?\s*(.*?)(?=Danh mục|Hóa đơn|Vận đơn|Tờ khai|Chủng loại)", text))
    invoice_no = _strip_status(_first(r"Hóa đơn số\s*:?\s*(.*?)(?=Vận đơn|Tờ khai|Chủng loại|Số/khối lượng)", text))
    bill_of_lading_no = _strip_status(_first(r"Vận đơn số\s*:?\s*(.*?)(?=Tờ khai|Chủng loại|Số/khối lượng)", text))
    deposit_confirmation = _clean_deposit_confirmation(_first(
        r"Giấy xác nhận đã ký quỹ bảo đảm phế liệu nhập khẩu số\s*:?\s*(.*?)(?=-\s*Thông tin|Thông tin về lô hàng|Hợp đồng)",
        text,
    ))
    customs_declaration_no = _strip_status(_first(r"Tờ khai hàng hóa nhập khẩu số\s*:?\s*(.*?)(?=Chủng loại|Số/khối lượng)", text))
    product_name = _strip_status(
        _first(r"Chủng loại phế liệu nhập khẩu\s*:?\s*(.*?)(?=Số/khối lượng|3\. Nội dung|Nội dung và kết quả)", text)
    )
    quantity = _first(r"Số/khối lượng hàng theo khai báo\s*:?\s*(.*?)(?=3\.|Nội dung|$)", text)
    if not quantity:
        quantity = _first(r"Số lượng Container\s*:?\s*([0-9]{1,3}\s*cont)", text)
    hs_code = _first(r"MÃ\s*HS\s*:?\s*([0-9]+)", text)
    container_count = (
        _first(r"(?:4\.1\.\s*)?Số lượng container\s*:?\s*([^\n\r.]+?\([^)]+\)?)", text)
        or _first(r"Số lượng container\s*:?\s*([^\n\r]+)", text)
        or _first(r"([0-9]{1,3}\s*x\s*40['’]\s*CONTAINER)", quantity)
    )
    container_count_short = container_count or _first(r"([0-9]{1,3})\s*x\s*40", quantity)
    if not container_count_short:
        container_count_short = quantity

    return CertificateData(
        certificate_no=certificate_no,
        certificate_date=certificate_date,
        customer_name=customer_name.replace(".", ""),
        customer_address=customer_address,
        inspection_location=inspection_location,
        contract_no=contract_no,
        invoice_no=invoice_no,
        bill_of_lading_no=bill_of_lading_no,
        deposit_confirmation=deposit_confirmation,
        customs_declaration_no=customs_declaration_no,
        product_name=product_name,
        hs_code=hs_code,
        quantity=quantity,
        container_count=container_count,
        container_count_short=container_count_short,
    )


def extract_certificate(path: Path) -> CertificateData:
    text = _read_docx_text(path)
    stop_labels = [
        "Địa chỉ",
        "Giấy xác nhận",
        "Giấy xác nhận ký quỹ",
        "Địa điểm kiểm tra, giám định",
        "Thời gian",
        "Hợp đồng số",
        "Danh mục hàng hóa",
        "Hóa đơn số",
        "Vận đơn số",
        "Tờ khai",
        "Chủng loại",
        "Số lượng hàng",
        "2. Nội dung",
    ]

    customer_name = _after_label(text, "Tên tổ chức, cá nhân", stop_labels)
    customer_address = _after_label(text, "Địa chỉ", stop_labels)
    inspection_location = _after_label(text, "Địa điểm kiểm tra, giám định", stop_labels)
    inspection_date = _after_label(text, "Thời gian kiểm tra, giám định", stop_labels)
    contract_no = _first(r"Hợp đồng số\s*:?\s*(.*?)\s*Ngày", text) or _after_label(text, "Hợp đồng số", stop_labels)
    invoice_no = _first(r"Hóa đơn số\s*:?\s*(.*?)\s*Ngày", text) or _after_label(text, "Hóa đơn số", stop_labels)
    bill_of_lading_no = _first(r"Vận đơn số\s*:?\s*(.*?)\s*Ngày", text) or _after_label(text, "Vận đơn số", stop_labels)
    deposit_confirmation = _clean_deposit_confirmation(_first(
        r"Giấy xác nhận ký quỹ đảm bảo\s*(.*?)\s*(?:phế liệu nhập khẩu số\s*)?(?=Địa điểm kiểm tra, giám định|Thời gian)",
        text,
    ))
    quantity = _after_label(text, "Số lượng hàng", stop_labels)
    product_name = _after_label(text, "Chủng loại phế liệu nhập khẩu", stop_labels)

    hs_code = _first(r"MÃ\s*HS\s*:?\s*([0-9]+)", text)
    certificate_no = _extract_cert_no(text)
    certificate_date = _extract_cert_date(text)
    customs_declaration_no = (
        _first(r"([0-9]{10,12})\s*Ngày\s+đăng\s+ký", text)
        or _first(r"Tờ\s+khai.*?số\s*:?\s*([0-9]{8,15}|[A-Z0-9.\-/]{5,30})", text)
    )
    container_count = (
        _first(r"(?:4\.1\.\s*)?Số lượng container\s*:?\s*([^\n\r.]+?\([^)]+\)?)", text)
        or _first(r"Số lượng container\s*:?\s*([^\n\r]+)", text)
        or _first(r"Số lượng hàng\s*:?\s*([^/\n\r]+(?:CONTAINER|CONT)?)", text)
    )
    container_count_short = container_count or _first(r"([0-9]{1,3}|x{1,3})\s*x\s*40", quantity)
    if not container_count_short:
        container_count_short = quantity

    return CertificateData(
        certificate_no=certificate_no,
        certificate_date=certificate_date,
        customer_name=customer_name.replace(".", ""),
        customer_address=customer_address,
        inspection_location=inspection_location,
        inspection_date=inspection_date,
        contract_no=contract_no,
        invoice_no=invoice_no,
        bill_of_lading_no=bill_of_lading_no,
        deposit_confirmation=deposit_confirmation,
        customs_declaration_no=customs_declaration_no,
        product_name=product_name,
        hs_code=hs_code,
        quantity=quantity,
        container_count=container_count,
        container_count_short=container_count_short,
    )
