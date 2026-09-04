from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document


def normalize_label(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip().lower()


def compact_label(text: str) -> str:
    return re.sub(r"[\W_]+", "", normalize_label(text), flags=re.UNICODE)


def is_placeholder_value(value: str) -> bool:
    if re.search(r"\b[xX]{2,}\b", value):
        return True
    compact = compact_label(value)
    if not compact:
        return False
    without_units = re.sub(r"(cont|container|kgm|kg|ngay|date)+", "", compact)
    return bool(without_units) and not re.sub(r"x+", "", without_units)


def certificate_date_long(value: str) -> str:
    match = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", value)
    if not match:
        return value
    day, month, year = match.groups()
    return f"ngày {int(day):02d} tháng {int(month):02d} năm {year}"


def convert_doc_to_docx(path: Path) -> Path:
    path = path.resolve()
    output = path.with_suffix(".docx")

    # 1. Try Windows MS Word COM Automation if on Windows
    if sys.platform == "win32":
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore

            pythoncom.CoInitialize()
            word = None
            try:
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(str(path))
                doc.SaveAs(str(output), FileFormat=16)
                doc.Close(False)
                return output
            finally:
                if word is not None:
                    word.Quit()
                pythoncom.CoUninitialize()
        except Exception:
            pass

    # 2. Try LibreOffice CLI (soffice / libreoffice --headless) for Linux / Docker
    for cmd in ["soffice", "libreoffice"]:
        try:
            import subprocess

            res = subprocess.run(
                [cmd, "--headless", "--convert-to", "docx", str(path), "--outdir", str(path.parent)],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                check=False,
            )
            if output.exists():
                return output
        except FileNotFoundError:
            continue

    raise RuntimeError("Không thể chuyển đổi file .doc. Cần cài đặt MS Word (Windows) hoặc LibreOffice (Linux Docker).")


def _replace_span_in_runs(runs, target_start: int, target_end: int, new_text: str) -> None:
    if not runs:
        return
    full_text = "".join(r.text for r in runs)
    if target_start > len(full_text) or target_end > len(full_text):
        target_start = 0
        target_end = len(full_text)
    curr_offset = 0
    first_replaced = False
    for r in runs:
        r_len = len(r.text)
        r_start = curr_offset
        r_end = curr_offset + r_len
        curr_offset = r_end

        if r_end <= target_start or r_start >= target_end:
            if target_start == target_end == r_end and not first_replaced:
                r.text = r.text + new_text
                first_replaced = True
            continue

        sub_start = max(0, target_start - r_start)
        sub_end = min(r_len, target_end - r_start)

        prefix = r.text[:sub_start]
        suffix = r.text[sub_end:]

        if not first_replaced:
            r.text = prefix + new_text + suffix
            first_replaced = True
        else:
            r.text = prefix + suffix


def _replace_paragraph_value(paragraph, label: str, value: str) -> bool:
    if not value:
        return False
    text = paragraph.text
    clean_label = label.rstrip(":").strip()
    if not clean_label:
        return False

    norm_label = normalize_label(clean_label)
    norm_text = normalize_label(text)

    is_date_label = any(k in norm_label for k in ["ngay", "thang"])
    is_date_text = any(k in norm_text for k in ["ngay", "thang"])

    if is_date_label and is_date_text:
        date_pattern = re.compile(
            r"((?:Hải\s+Phòng|Hà\s+Nội|TP\.?\s*HCM|TP\.?\s*Hồ\s+Chí\s+Minh)?\s*,?\s*ngày\s+)[^\t\n\r]*|(ngày\s+.*?tháng\s+.*?năm\s+.*?)(?=\t|\n|\r|$)",
            re.IGNORECASE,
        )
        m = date_pattern.search(text)
        if m:
            if m.group(1):
                _replace_span_in_runs(paragraph.runs, m.start(0) + len(m.group(1)), m.end(0), value)
            else:
                _replace_span_in_runs(paragraph.runs, m.start(0), m.end(0), value)
            return True

    has_colon = label.rstrip().endswith(":")
    colon_part = r":\s*" if has_colon else r"\s*:?\s*"
    label_pattern = re.compile(rf"((?:^|\s|\t|\b){re.escape(clean_label)}{colon_part})([^\t\n\r]*)", flags=re.IGNORECASE)
    m = label_pattern.search(text)
    if m:
        _replace_span_in_runs(paragraph.runs, m.start(2), m.end(2), value)
        return True

    return False


def _replace_paragraph_regex(paragraph, pattern: str, value: str) -> bool:
    if not value:
        return False
    text = paragraph.text
    try:
        pat = re.compile(pattern, flags=re.IGNORECASE)
    except Exception:
        return False
    m = pat.search(text)
    if not m:
        return False
    _replace_span_in_runs(paragraph.runs, m.start(0), m.end(0), value)
    return True


def _replace_table_value(table, label: str, value: str) -> bool:
    if not value:
        return False
    wanted = normalize_label(label).rstrip(":")
    changed = False
    for row in table.rows:
        cells = row.cells
        for index, cell in enumerate(cells):
            cell_text = normalize_label(cell.text).rstrip(":")
            if wanted in cell_text or compact_label(label) in compact_label(cell.text):
                target = cells[index + 1] if index + 1 < len(cells) else cell
                if target is cell:
                    for paragraph in cell.paragraphs:
                        changed = _replace_paragraph_value(paragraph, label, value) or changed
                else:
                    if target.paragraphs:
                        p = target.paragraphs[0]
                        if p.runs:
                            _replace_span_in_runs(p.runs, 0, len(p.text), value)
                        else:
                            p.text = value
                        for extra_p in target.paragraphs[1:]:
                            for r in extra_p.runs:
                                r.text = ""
                    else:
                        target.text = value
                    changed = True
    return changed



def update_labels(update: dict) -> list[str]:
    labels = update.get("labels")
    if labels:
        return labels
    label = update.get("label")
    return [label] if label else []


def _all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for section in doc.sections:
        for container in [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]:
            for paragraph in container.paragraphs:
                yield paragraph


def _body_paragraphs(doc):
    return list(doc.paragraphs)


def _all_tables(doc):
    for table in doc.tables:
        yield table
    for section in doc.sections:
        for container in [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]:
            for table in container.tables:
                yield table


def update_docx(template_path: Path, output_path: Path, updates: list[dict], data: dict[str, str]) -> list[str]:
    doc = Document(str(template_path))
    missing: list[str] = []

    for update in updates:
        labels = update_labels(update)
        label = labels[0] if labels else update.get("pattern", "")
        value = data.get(update["field"], "")
        if update.get("format") == "certificate_date_long":
            value = certificate_date_long(value)
        if is_placeholder_value(value):
            missing.append(f"{label} (du lieu nguon dang la {value})")
            continue
        changed = False
        if update.get("mode") == "regex":
            for paragraph in _all_paragraphs(doc):
                changed = _replace_paragraph_regex(paragraph, update["pattern"], value) or changed
        else:
            for candidate_label in labels:
                for table in _all_tables(doc):
                    changed = _replace_table_value(table, candidate_label, value) or changed
                for paragraph in _all_paragraphs(doc):
                    changed = _replace_paragraph_value(paragraph, candidate_label, value) or changed
                if changed:
                    break
            clear_next = update.get("clear_next_if_starts")
            if clear_next and changed:
                paragraphs = _body_paragraphs(doc)
                for index, paragraph in enumerate(paragraphs[:-1]):
                    if any(compact_label(item) in compact_label(paragraph.text) for item in labels):
                        next_paragraph = paragraphs[index + 1]
                        if compact_label(next_paragraph.text).startswith(compact_label(clear_next)):
                            for run in next_paragraph.runs:
                                run.text = ""
        if not changed:
            missing.append(label)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return missing


def load_client_config(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def prepare_template(path: Path) -> Path:
    path = path.resolve()
    if path.suffix.lower() == ".doc":
        docx_path = path.with_suffix(".docx")
        if docx_path.exists():
            return docx_path
        return convert_doc_to_docx(path)
    return path


def copy_sample_templates(sample_dir: Path, project_templates: Path) -> None:
    project_templates.mkdir(parents=True, exist_ok=True)
    for source in sample_dir.iterdir():
        suffix = source.suffix.lower()
        if suffix not in {".doc", ".docx"}:
            continue
        target_name = {
            "phiếu changxin.docx": "phieu_CHANGXIN.docx",
            "17020icc-i-qt04-bbht_klm.doc": "bien_ban_CHANGXIN.doc",
        }.get(source.name.lower())
        if target_name:
            shutil.copy2(source, project_templates / target_name)
