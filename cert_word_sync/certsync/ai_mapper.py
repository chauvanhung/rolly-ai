from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any

import requests
from docx import Document

from certsync.extractor import CertificateData, extract_certificate
from certsync.client_config import make_template_config
from certsync.updater import prepare_template


def _read_full_docx_text(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    lines: list[str] = []

    # 1. Header
    for sec in doc.sections:
        for container_name, container in (("HEADER", sec.header), ("FIRST HEADER", sec.first_page_header), ("EVEN HEADER", sec.even_page_header)):
            if not container:
                continue
            for p in container.paragraphs:
                if p.text.strip():
                    lines.append(f"[{container_name}] {p.text.strip()}")
            for t in container.tables:
                for r in t.rows:
                    row_txt = " | ".join(c.text.strip() for c in r.cells if c.text.strip())
                    if row_txt:
                        lines.append(f"[{container_name} TABLE] {row_txt}")

    # 2. Body Paragraphs & Tables
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for t in doc.tables:
        for r in t.rows:
            row_txt = " | ".join(c.text.strip() for c in r.cells if c.text.strip())
            if row_txt:
                lines.append(f"[TABLE ROW] {row_txt}")

    # 3. Text Boxes
    try:
        tb_elements = doc._element.xpath(".//*[local-name()='txbxContent']")
        for tb in tb_elements:
            txt = "".join(node.text for node in tb.xpath(".//*[local-name()='t']") if node.text)
            if txt.strip():
                lines.append(f"[TEXT BOX] {txt.strip()}")
    except Exception:
        pass

    # 4. Footers: official forms may place fields here.
    for sec in doc.sections:
        for container_name, container in (("FOOTER", sec.footer), ("FIRST FOOTER", sec.first_page_footer), ("EVEN FOOTER", sec.even_page_footer)):
            if not container:
                continue
            for p in container.paragraphs:
                if p.text.strip():
                    lines.append(f"[{container_name}] {p.text.strip()}")
            for t in container.tables:
                for r in t.rows:
                    row_txt = " | ".join(c.text.strip() for c in r.cells if c.text.strip())
                    if row_txt:
                        lines.append(f"[{container_name} TABLE] {row_txt}")

    return "\n".join(lines)


def call_ai_chat_completion(prompt: str, system_instruction: str = "") -> str | None:
    api_key = (
        os.environ.get("GEMINI_API_KEY")
        or os.environ.get("OPENAI_API_KEY")
        or os.environ.get("AI_API_KEY")
        or ""
    ).strip()

    if not api_key:
        return None

    # Google Gemini OpenAI-Compatible Endpoint or Standard OpenAI Endpoint
    openai_base_url = (os.environ.get("OPENAI_BASE_URL") or os.environ.get("AI_BASE_URL") or "").strip()
    if not openai_base_url:
        if api_key.startswith("AIza"):
            openai_base_url = "https://generativelanguage.googleapis.com/v1beta/openai/"
        else:
            openai_base_url = "https://api.openai.com/v1"

    model_name = (
        os.environ.get("OPENAI_LLM_MODEL")
        or os.environ.get("AI_MODEL")
        or ("gemini-2.5-flash" if "generativelanguage" in openai_base_url else "dev")
    )
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_key}",
    }
    messages = []
    if system_instruction:
        messages.append({"role": "system", "content": system_instruction})
    messages.append({"role": "user", "content": prompt})

    try:
        url = f"{openai_base_url.rstrip('/')}/chat/completions"
        payload = {
            "model": model_name,
            "messages": messages,
            "temperature": 0.1,
            "response_format": {"type": "json_object"},
        }
        res = requests.post(url, headers=headers, json=payload, timeout=20)
        if res.status_code != 200:
            payload.pop("response_format", None)
            res = requests.post(url, headers=headers, json=payload, timeout=20)
        if res.status_code == 200:
            data = res.json()
            return data["choices"][0]["message"]["content"]
    except Exception as e:
        print(f"[AI Mapper] AI API Call Warning: {str(e)}")

    return None


def _normalize_template_config(config: dict[str, Any], client_id: str, phieu_path: Path, bien_ban_path: Path) -> dict[str, Any]:
    templates = config.get("templates") if isinstance(config, dict) else None
    if not isinstance(templates, list):
        templates = []
    expected = [(phieu_path.name, f"phieu_{client_id}_da_cap_nhat.docx"), (bien_ban_path.name, f"bien_ban_{client_id}_da_cap_nhat.docx")]
    normalized = []
    for index, (source, output) in enumerate(expected):
        item = templates[index] if index < len(templates) and isinstance(templates[index], dict) else {}
        updates = item.get("updates", [])
        if not isinstance(updates, list):
            updates = []
        clean_updates = []
        for update in updates:
            if not isinstance(update, dict) or not update.get("field"):
                continue
            labels = update.get("labels") or ([update["label"]] if update.get("label") else [])
            if isinstance(labels, str):
                labels = [labels]
            labels = [str(label).strip() for label in labels if str(label).strip()]
            if labels:
                clean_updates.append({**update, "labels": labels})
        normalized.append({"source": str(item.get("source") or source), "output": str(item.get("output") or output), "updates": clean_updates})
    return {"name": str(config.get("name") or client_id.upper()), "templates": normalized}


def _clean_updates(updates: Any) -> list[dict[str, Any]]:
    if not isinstance(updates, list):
        return []
    clean_updates = []
    for update in updates:
        if not isinstance(update, dict) or not update.get("field"):
            continue
        labels = update.get("labels") or ([update["label"]] if update.get("label") else [])
        if isinstance(labels, str):
            labels = [labels]
        labels = [str(label).strip() for label in labels if str(label).strip()]
        if labels:
            clean_updates.append({**update, "labels": labels})
    return clean_updates


def _normalize_single_template_config(config: dict[str, Any], client_id: str, template_path: Path, kind: str) -> dict[str, Any]:
    template = make_template_config(client_id, kind, template_path.name)
    candidates = []
    if isinstance(config, dict):
        if isinstance(config.get("templates"), list):
            candidates = [item for item in config["templates"] if isinstance(item, dict)]
        elif isinstance(config.get("updates"), list):
            candidates = [config]
    if candidates:
        updates = _clean_updates(candidates[0].get("updates", []))
        if updates:
            template["updates"] = updates
    return {"name": str(config.get("name") or client_id.upper()) if isinstance(config, dict) else client_id.upper(), "templates": [template]}


def extract_cert_with_ai(docx_path: Path) -> CertificateData:
    full_text = _read_full_docx_text(docx_path)
    fallback_data = extract_certificate(docx_path)

    system_prompt = (
        "Bạn là Chuyên gia AI Giám định Văn bản & Bóc tách Dữ liệu Chứng thư. "
        "Hãy đọc văn bản chứng thư được cung cấp và trích xuất đúng các thông tin theo cấu trúc JSON. "
        "Nếu không có thông tin trong văn bản, hãy để chuỗi rỗng \"\"."
    )
    user_prompt = f"""
Hãy bóc tách các trường thông tin sau dưới dạng JSON duy nhất với các khóa:
- certificate_no: Số chứng thư (Ví dụ: ICC/08.26/66001.22)
- certificate_date: Ngày cấp chứng thư (Ví dụ: ngày 30 tháng 08 năm 2026)
- customer_name: Tên công ty / tổ chức nhập khẩu (Ví dụ: CÔNG TY TNHH CHANG XIN (VIỆT NAM))
- customer_address: Địa chỉ công ty nhập khẩu
- inspection_location: Địa điểm kiểm tra / giám định
- contract_no: Số hợp đồng
- invoice_no: Số hóa đơn / Số Invoice
- bill_of_lading_no: Số vận đơn / Bill of Lading
- deposit_confirmation: Số giấy xác nhận ký quỹ
- customs_declaration_no: Số tờ khai hải quan (Ví dụ: 10855545725)
- product_name: Tên hàng hóa / phế liệu
- hs_code: Mã HS Code
- container_count: Chi tiết số lượng container kèm theo danh sách số container trong ngoặc (Ví dụ: 03 x 40’container (MRSU3689815; HASU4500633; CIPU5005173))
- container_count_short: Số lượng container kèm danh sách số container trong ngoặc (Ví dụ: 03 x 40’container (MRSU3689815; HASU4500633; CIPU5005173))

NỘI DUNG VĂN BẢN CHỨNG THƯ:
---
{full_text}
---
    """

    ai_json_str = call_ai_chat_completion(user_prompt, system_prompt)
    if not ai_json_str:
        return fallback_data

    try:
        parsed = json.loads(ai_json_str)
        return CertificateData(
            certificate_no=parsed.get("certificate_no") or fallback_data.certificate_no,
            certificate_date=parsed.get("certificate_date") or fallback_data.certificate_date,
            customer_name=parsed.get("customer_name") or fallback_data.customer_name,
            customer_address=parsed.get("customer_address") or fallback_data.customer_address,
            inspection_location=parsed.get("inspection_location") or fallback_data.inspection_location,
            contract_no=parsed.get("contract_no") or fallback_data.contract_no,
            invoice_no=parsed.get("invoice_no") or fallback_data.invoice_no,
            bill_of_lading_no=parsed.get("bill_of_lading_no") or fallback_data.bill_of_lading_no,
            deposit_confirmation=parsed.get("deposit_confirmation") or fallback_data.deposit_confirmation,
            customs_declaration_no=parsed.get("customs_declaration_no") or fallback_data.customs_declaration_no,
            product_name=parsed.get("product_name") or fallback_data.product_name,
            hs_code=parsed.get("hs_code") or fallback_data.hs_code,
            container_count=parsed.get("container_count") or fallback_data.container_count,
            container_count_short=parsed.get("container_count_short") or fallback_data.container_count_short,
        )
    except Exception as e:
        print(f"[AI Mapper] AI JSON parsing failed: {e}")
        return fallback_data


def ai_generate_template_config(phieu_path: Path, bien_ban_path: Path, client_id: str) -> dict:
    phieu_text = _read_full_docx_text(phieu_path)
    bien_ban_text = _read_full_docx_text(bien_ban_path)

    system_prompt = (
        "Bạn là Kỹ sư AI Cấu hình Mẫu Word. Hãy phân tích nội dung hai file mẫu 'Phiếu xác nhận' và 'Biên bản kiểm tra' "
        "và trả về cấu hình JSON danh sách các trường cần cập nhật tự động."
    )
    user_prompt = f"""
Dưới đây là văn bản 2 file mẫu Word của khách hàng '{client_id}':

LƯU Ý: Đối với các nhãn liên quan đến 'Số lượng Container' hoặc 'Số container', LUÔN ƯU TIÊN ghép vào field 'container_count' (chứa đầy đủ số container kèm ngoặc đơn). Chỉ dùng 'container_count_short' nếu nhãn yêu cầu rõ ràng rút gọn.

=== FILE 1: PHIẾU XÁC NHẬN ===
{phieu_text}

=== FILE 2: BIÊN BẢN KIỂM TRA ===
{bien_ban_text}

Hãy trả về JSON cấu hình chuẩn theo định dạng sau:
{{
  "name": "{client_id.upper()}",
  "templates": [
    {{
      "source": "{phieu_path.name}",
      "output": "phieu_{client_id}_da_cap_nhat.docx",
      "updates": [
        {{ "labels": ["Số chứng thư", "Mục số"], "field": "certificate_no" }},
        {{ "labels": ["Tên tổ chức", "Tên đơn vị"], "field": "customer_name" }},
        {{ "labels": ["Số tờ khai", "Tờ khai hải quan"], "field": "customs_declaration_no" }}
      ]
    }},
    {{
      "source": "{bien_ban_path.name}",
      "output": "bien_ban_{client_id}_da_cap_nhat.docx",
      "updates": [
        {{ "labels": ["Số:", "Số chứng thư"], "field": "certificate_no" }},
        {{ "labels": ["Tên tổ chức", "Đơn vị nhập khẩu"], "field": "customer_name" }}
      ]
    }}
  ]
}}
    """

    ai_json_str = call_ai_chat_completion(user_prompt, system_prompt)
    if ai_json_str:
        try:
            config = _normalize_template_config(json.loads(ai_json_str), client_id, phieu_path, bien_ban_path)
            config["mapping_source"] = "ai"
            return config
        except Exception:
            pass

    # Default structural template config
    return {
        "name": client_id.upper(),
        "mapping_source": "default",
        "templates": [
            {
                "source": phieu_path.name,
                "output": f"phieu_{client_id}_da_cap_nhat.docx",
                "updates": [
                    {"labels": ["Số chứng thư", "Mục số"], "field": "certificate_no"},
                    {"labels": ["Tên tổ chức", "Tên đơn vị"], "field": "customer_name"},
                    {"labels": ["Địa chỉ"], "field": "customer_address"},
                    {"labels": ["Số tờ khai"], "field": "customs_declaration_no"},
                    {"labels": ["Số hợp đồng"], "field": "contract_no"},
                ],
            },
            {
                "source": bien_ban_path.name,
                "output": f"bien_ban_{client_id}_da_cap_nhat.docx",
                "updates": [
                    {"labels": ["Số:", "Số chứng thư"], "field": "certificate_no"},
                    {"labels": ["Tên tổ chức"], "field": "customer_name"},
                    {"labels": ["Địa chỉ"], "field": "customer_address"},
                    {"labels": ["Tờ khai"], "field": "customs_declaration_no"},
                ],
            },
        ],
    }


def ai_generate_single_template_config(template_path: Path, client_id: str, kind: str) -> dict:
    template_text = _read_full_docx_text(template_path)
    kind_label = "Phiếu xác nhận" if kind == "phieu" else "Biên bản kiểm tra"

    system_prompt = (
        "Bạn là kỹ sư cấu hình mẫu Word. Hãy đọc một file mẫu Word và trả về JSON mapping "
        "các nhãn trong form sang field dữ liệu chuẩn. Chỉ trả JSON hợp lệ."
    )
    user_prompt = f"""
Đây là nội dung file mẫu '{kind_label}' của khách hàng '{client_id}'.

Nhiệm vụ:
- Tìm các nhãn trong form mà cần được tự động điền.
- Ghép từng nhãn vào đúng field chuẩn.
- Chỉ dùng field trong danh sách này:
  certificate_no, certificate_date, customer_name, customer_address,
  inspection_location, deposit_confirmation, contract_no, invoice_no,
  bill_of_lading_no, customs_declaration_no, product_name, quantity,
  container_count, container_count_short, hs_code
- QUY TẮC CONTAINER: Đối với nhãn về 'Số lượng Container' hoặc 'Số container', LUÔN ƯU TIÊN chọn field 'container_count' (bóc tách đầy đủ kèm số container trong ngoặc). Chỉ dùng 'container_count_short' khi nhãn mẫu có yêu cầu rút gọn rõ ràng.
- Trả về JSON theo dạng:
{{
  "name": "{client_id.upper()}",
  "templates": [
    {{
      "source": "{template_path.name}",
      "updates": [
        {{ "labels": ["Số chứng thư", "Số:"], "field": "certificate_no" }}
      ]
    }}
  ]
}}

NỘI DUNG FILE MẪU:
---
{template_text}
---
    """

    ai_json_str = call_ai_chat_completion(user_prompt, system_prompt)
    if ai_json_str:
        try:
            config = _normalize_single_template_config(json.loads(ai_json_str), client_id, template_path, kind)
            config["mapping_source"] = "ai"
            return config
        except Exception:
            pass
    return {"name": client_id.upper(), "mapping_source": "default", "templates": [make_template_config(client_id, kind, template_path.name)]}
