from __future__ import annotations

import html
import json
import os
import shutil
import secrets
import tempfile
from datetime import timedelta
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from flask import (
    Flask,
    jsonify,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
    url_for,
)
from werkzeug.utils import secure_filename
import mammoth

from certsync.auth import (
    ADMIN_EMAIL,
    admin_required,
    admin_reset_password,
    authenticate_user,
    create_user,
    delete_user,
    exchange_google_code_and_get_user,
    get_google_auth_url,
    get_user_by_id,
    init_auth_db,
    list_users,
    login_required,
    verify_google_id_token_and_get_user,
)
from certsync.client_config import (
    build_client_config,
    make_template_config,
    normalize_client_id,
    save_client_config,
)
from certsync.extractor import extract_certificate
from certsync.updater import load_client_config, prepare_template, update_docx

BASE_DIR = Path(__file__).resolve().parent
CLIENTS_DIR = BASE_DIR / "clients"
TEMPLATES_DIR = BASE_DIR / "templates"
OUTPUT_DIR = BASE_DIR / "outputs"
AUTH_DB_PATH = BASE_DIR / "data" / "auth.sqlite3"

FIELD_LABELS_VI = {
    "certificate_no": "Số chứng thư",
    "certificate_date": "Ngày cấp chứng thư",
    "customer_name": "Tên khách hàng / tổ chức",
    "customer_address": "Địa chỉ khách hàng",
    "inspection_location": "Địa điểm giám định / lấy mẫu",
    "deposit_confirmation": "Số giấy xác nhận ký quỹ",
    "contract_no": "Số hợp đồng",
    "invoice_no": "Số hóa đơn",
    "bill_of_lading_no": "Số vận đơn",
    "customs_declaration_no": "Số tờ khai hải quan",
    "product_name": "Chủng loại hàng / sản phẩm",
    "quantity": "Số lượng / khối lượng hàng",
    "container_count": "Số lượng container (Đầy đủ kèm ngoặc)",
    "container_count_short": "Số lượng container (Rút gọn)",
    "hs_code": "Mã HS Code",
}

load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "cert_word_sync_secret_key_2026_super_secure")
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)

init_auth_db(AUTH_DB_PATH)


def convert_docx_to_full_html(file_path: Path) -> str:
    doc = Document(str(file_path))
    html_parts: list[str] = []

    # 1. Include Header paragraphs and tables
    for sec in doc.sections:
        for container in (sec.header, sec.first_page_header, sec.even_page_header):
            if not container:
                continue
            for p in container.paragraphs:
                if p.text.strip():
                    html_parts.append(f"<p><strong>[Header]</strong> {html.escape(p.text.strip())}</p>")
            for t in container.tables:
                html_parts.append("<table border='1' style='border-collapse:collapse;width:100%;margin-bottom:10px;'>")
                for r in t.rows:
                    html_parts.append("<tr>")
                    for c in r.cells:
                        html_parts.append(f"<td style='padding:4px;'>{html.escape(c.text.strip())}</td>")
                    html_parts.append("</tr>")
                html_parts.append("</table>")

    # 2. Main Body Content via Mammoth
    with open(file_path, "rb") as docx_file:
        m_result = mammoth.convert_to_html(docx_file)
        html_parts.append(m_result.value)

    return "".join(html_parts)


def client_ids() -> list[str]:
    return sorted(path.stem for path in CLIENTS_DIR.glob("*.json"))


def client_details() -> list[dict]:
    results = []
    for path in sorted(CLIENTS_DIR.glob("*.json")):
        client_id = path.stem
        try:
            cfg = load_client_config(path)
            templates = cfg.get("templates", [])
            template_names = [t.get("source") for t in templates if t.get("source")]
            total_updates = sum(len(t.get("updates", [])) for t in templates)
            results.append({
                "id": client_id,
                "name": cfg.get("name", client_id.upper()),
                "templates_count": len(templates),
                "template_names": template_names,
                "total_updates": total_updates,
                "mapping_source": cfg.get("mapping_source", "manual"),
            })
        except Exception:
            results.append({
                "id": client_id,
                "name": client_id.upper(),
                "templates_count": 0,
                "template_names": [],
                "total_updates": 0,
                "mapping_source": "error",
            })
    return results


def get_client_config_path(client_id: str) -> Path:
    safe_client_id = normalize_client_id(client_id)
    return CLIENTS_DIR / f"{safe_client_id}.json"


def assert_inside(child: Path, parent: Path) -> Path:
    child = child.resolve()
    parent = parent.resolve()
    if child != parent and parent not in child.parents:
        raise ValueError(f"Unsafe path outside project: {child}")
    return child


def is_word_upload(uploaded) -> bool:
    return bool(uploaded and uploaded.filename and uploaded.filename.lower().endswith((".doc", ".docx")))


def saved_upload_name(uploaded, fallback: str) -> str:
    filename = secure_filename(uploaded.filename or "")
    return filename or fallback


def template_kind(template: dict) -> str | None:
    text = f"{template.get('source', '')} {template.get('output', '')}".lower()
    if "bien_ban" in text or "bien-ban" in text or "bien" in text:
        return "bien_ban"
    if "phieu" in text or "phieu-xac-nhan" in text:
        return "phieu"
    return None


def has_template_kind(config: dict, kind: str) -> bool:
    return any(template_kind(template) == kind for template in config.get("templates", []))


def set_user_session(user: dict) -> None:
    session.permanent = True
    session["user_id"] = user["id"]
    session["user_email"] = user["email"]
    session["user_name"] = user["full_name"]


@app.context_processor
def inject_user():
    user_id = session.get("user_id")
    current_user = get_user_by_id(AUTH_DB_PATH, user_id) if user_id else None
    return {"current_user": current_user}


# --- Authentication Routes ---


@app.route("/login", methods=["GET"])
def login_page():
    if session.get("user_id"):
        return redirect(url_for("index"))

    google_client_id = (os.environ.get("GOOGLE_CLIENT_ID") or "").strip()
    google_auth_url = ""
    if google_client_id:
        redirect_uri = get_google_redirect_uri()
        oauth_state = secrets.token_urlsafe(16)
        session["oauth_state"] = oauth_state
        google_auth_url = get_google_auth_url(google_client_id, redirect_uri, oauth_state)

    error = request.args.get("error")
    msg = request.args.get("msg")
    return render_template(
        "login.html",
        error=error,
        msg=msg,
        google_auth_url=google_auth_url,
        google_client_id=google_client_id,
    )


@app.route("/login", methods=["POST"])
def login_post():
    email = request.form.get("email", "")
    password = request.form.get("password", "")
    user = authenticate_user(AUTH_DB_PATH, email, password)
    if not user:
        google_client_id = (os.environ.get("GOOGLE_CLIENT_ID") or "").strip()
        google_auth_url = ""
        if google_client_id:
            redirect_uri = get_google_redirect_uri()
            google_auth_url = get_google_auth_url(google_client_id, redirect_uri)
        return render_template(
            "login.html",
            error="Email hoặc mật khẩu không chính xác.",
            google_auth_url=google_auth_url,
        )

    set_user_session(user)
    return redirect(url_for("index"))


@app.route("/register", methods=["GET"])
def register_page():
    if session.get("user_id"):
        return redirect(url_for("index"))
    return render_template("register.html")


@app.route("/register", methods=["POST"])
def register_post():
    email = request.form.get("email", "")
    password = request.form.get("password", "")
    full_name = request.form.get("full_name", "")

    ok, msg = create_user(AUTH_DB_PATH, email, password, full_name)
    if not ok:
        return render_template("register.html", error=msg)

    # Auto login on registration
    user = authenticate_user(AUTH_DB_PATH, email, password)
    if user:
        set_user_session(user)
        return redirect(url_for("index"))

    return redirect(url_for("login_page", msg=msg))


def get_google_redirect_uri() -> str:
    host = request.headers.get("X-Forwarded-Host") or request.host
    proto = request.headers.get("X-Forwarded-Proto") or ("https" if "rollyhub.com" in host else request.scheme)
    return f"{proto}://{host}/auth/google/callback"


@app.route("/auth/google", methods=["GET"])
def google_auth():
    google_client_id = (os.environ.get("GOOGLE_CLIENT_ID") or "").strip()
    if not google_client_id:
        return redirect(url_for("login_page", error="Google Auth chưa được cấu hình. Thiếu GOOGLE_CLIENT_ID."))

    redirect_uri = get_google_redirect_uri()
    oauth_state = secrets.token_urlsafe(16)
    session["oauth_state"] = oauth_state
    auth_url = get_google_auth_url(google_client_id, redirect_uri, oauth_state)
    return redirect(auth_url)


@app.route("/auth/google/callback", methods=["GET"])
def google_callback():
    code = request.args.get("code")
    state = request.args.get("state")

    if not code:
        return redirect(url_for("login_page", error="Đăng nhập Google bị hủy hoặc thất bại."))

    stored_state = session.pop("oauth_state", None)
    if stored_state and state and stored_state != state:
        return redirect(url_for("login_page", error="Xác thực OAuth state không hợp lệ."))

    google_client_id = (os.environ.get("GOOGLE_CLIENT_ID") or "").strip()
    google_client_secret = (os.environ.get("GOOGLE_CLIENT_SECRET") or "").strip()
    redirect_uri = get_google_redirect_uri()

    if not google_client_id or not google_client_secret:
        return redirect(url_for("login_page", error="Thiếu cấu hình GOOGLE_CLIENT_ID hoặc GOOGLE_CLIENT_SECRET trên server."))

    user_dict, msg = exchange_google_code_and_get_user(
        db_path=AUTH_DB_PATH,
        client_id=google_client_id,
        client_secret=google_client_secret,
        code=code,
        redirect_uri=redirect_uri,
    )

    if not user_dict:
        return redirect(url_for("login_page", error=msg))

    set_user_session(user_dict)
    return redirect(url_for("index"))


@app.route("/auth/google/idtoken", methods=["POST"])
def google_idtoken():
    data = request.get_json(silent=True) or {}
    id_token = data.get("credential") or request.form.get("credential")
    if not id_token:
        return jsonify({"ok": False, "error": "Thiếu Google credential token."}), 400

    google_client_id = (os.environ.get("GOOGLE_CLIENT_ID") or "").strip()
    user_dict, msg = verify_google_id_token_and_get_user(
        db_path=AUTH_DB_PATH, client_id=google_client_id, id_token=id_token
    )
    if not user_dict:
        return jsonify({"ok": False, "error": msg}), 400

    set_user_session(user_dict)
    return jsonify({"ok": True, "redirect": url_for("index")})


@app.route("/logout", methods=["GET", "POST"])
def logout():
    session.clear()
    return redirect(url_for("login_page"))


# --- Admin User Management Routes ---


@app.route("/admin/users", methods=["GET"])
@admin_required
def admin_users_page():
    users = list_users(AUTH_DB_PATH)
    msg = request.args.get("msg")
    error = request.args.get("error")
    return render_template("admin_users.html", users=users, msg=msg, error=error, system_settings=load_system_settings())


@app.route("/admin/users/create", methods=["POST"])
@admin_required
def admin_users_create():
    email = request.form.get("email", "")
    password = request.form.get("password", "")
    full_name = request.form.get("full_name", "")
    ok, msg = create_user(AUTH_DB_PATH, email, password, full_name)
    if not ok:
        users = list_users(AUTH_DB_PATH)
        return render_template("admin_users.html", users=users, error=msg)
    return redirect(url_for("admin_users_page", msg=f"Đã tạo tài khoản {email} thành công!"))


@app.route("/admin/users/<int:user_id>/reset_password", methods=["POST"])
@admin_required
def admin_users_reset_password(user_id: int):
    new_password = request.form.get("new_password", "")
    ok, msg = admin_reset_password(AUTH_DB_PATH, user_id, new_password)
    if not ok:
        return redirect(url_for("admin_users_page", error=msg))
    return redirect(url_for("admin_users_page", msg=msg))


@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@admin_required
def admin_users_delete(user_id: int):
    current_admin_id = session.get("user_id")
    ok, msg = delete_user(AUTH_DB_PATH, user_id, current_admin_id=current_admin_id)
    if not ok:
        return redirect(url_for("admin_users_page", error=msg))
    return redirect(url_for("admin_users_page", msg=msg))


from certsync.settings import load_system_settings, save_system_settings


@app.route("/admin/settings", methods=["POST"])
@admin_required
def update_admin_settings():
    enable_ai_extraction = request.form.get("enable_ai_extraction") == "1"
    enable_ai_template_learning = request.form.get("enable_ai_template_learning") == "1"
    save_system_settings({
        "enable_ai_extraction": enable_ai_extraction,
        "enable_ai_template_learning": enable_ai_template_learning,
    })
    return redirect(request.referrer or url_for("index"))


# --- App Main Routes ---


@app.route("/", methods=["GET"])
@login_required
def index():
    return render_template(
        "index.html",
        clients=client_ids(),
        client_list=client_details(),
        created=request.args.get("created"),
        error=request.args.get("error"),
        system_settings=load_system_settings(),
    )


from certsync.ai_mapper import ai_generate_single_template_config, ai_generate_template_config, extract_cert_with_ai


@app.route("/clients", methods=["GET", "POST"])
@admin_required
def create_client():
    if request.method == "GET":
        return redirect(url_for("index"))

    raw_client_id = request.form.get("client_id", "")
    if not raw_client_id.strip():
        return render_template("index.html", clients=client_ids(), client_list=client_details(), error="Vui lòng nhập Mã khách hàng."), 400

    client_id = normalize_client_id(raw_client_id)
    phieu = request.files.get("phieu_template")
    bien_ban = request.files.get("bien_ban_template")

    phieu_uploaded = bool(phieu and phieu.filename)
    bien_ban_uploaded = bool(bien_ban and bien_ban.filename)
    if not phieu_uploaded and not bien_ban_uploaded:
        return render_template("index.html", clients=client_ids(), client_list=client_details(), error="Cần tải lên ít nhất 1 file mẫu: Phiếu xác nhận hoặc Biên bản."), 400
    for uploaded in [phieu, bien_ban]:
        if uploaded and uploaded.filename and not is_word_upload(uploaded):
            return render_template("index.html", clients=client_ids(), client_list=client_details(), error="Template chỉ nhận file .doc hoặc .docx."), 400

    template_dir = TEMPLATES_DIR / client_id
    template_dir.mkdir(parents=True, exist_ok=True)
    phieu_name = None
    bien_ban_name = None
    phieu_path = None
    bien_ban_path = None
    if phieu_uploaded:
        phieu_name = saved_upload_name(phieu, f"phieu_{client_id}.docx")
        phieu_path = template_dir / phieu_name
        phieu.save(phieu_path)
    if bien_ban_uploaded:
        bien_ban_name = saved_upload_name(bien_ban, f"bien_ban_{client_id}.docx")
        bien_ban_path = template_dir / bien_ban_name
        bien_ban.save(bien_ban_path)

    system_settings = load_system_settings()
    use_ai = system_settings.get("enable_ai_template_learning", True)
    if "use_ai" in request.form:
        use_ai = request.form.get("use_ai") == "1"

    config = None
    if use_ai:
        if phieu_path and bien_ban_path:
            try:
                # Analyze converted DOCX content when a legacy .DOC file was uploaded.
                analysis_phieu_path = prepare_template(phieu_path)
                analysis_bien_ban_path = prepare_template(bien_ban_path)
                config = ai_generate_template_config(analysis_phieu_path, analysis_bien_ban_path, client_id)
                config["name"] = raw_client_id.strip()
                # Preserve uploaded source names so generation keeps the existing .DOC conversion path.
                for template, original_path in zip(config.get("templates", []), [phieu_path, bien_ban_path]):
                    template["source"] = original_path.name
            except Exception:
                config = None
        if config is None:
            if phieu_path:
                config = ai_generate_single_template_config(prepare_template(phieu_path), client_id, "phieu")
                config["templates"][0]["source"] = phieu_path.name
            elif bien_ban_path:
                config = ai_generate_single_template_config(prepare_template(bien_ban_path), client_id, "bien_ban")
                config["templates"][0]["source"] = bien_ban_path.name

    if config is None:
        config = build_client_config(client_id, phieu_name, bien_ban_name, raw_client_id.strip())
        config["mapping_source"] = "default"

    config["name"] = raw_client_id.strip()
    save_client_config(CLIENTS_DIR / f"{client_id}.json", config)
    return redirect(url_for("index", created=client_id))


@app.route("/clients/<client_id>/edit", methods=["GET"])
@admin_required
def edit_client_page(client_id: str):
    config_path = get_client_config_path(client_id)
    if not config_path.exists():
        return "Khách hàng không tồn tại.", 404
    config = load_client_config(config_path)
    template_dir = TEMPLATES_DIR / normalize_client_id(client_id)
    return render_template(
        "edit_client.html",
        client_id=normalize_client_id(client_id),
        config=config,
        config_json=json.dumps(config, ensure_ascii=False, indent=2),
        template_dir=template_dir,
        has_phieu=has_template_kind(config, "phieu"),
        has_bien_ban=has_template_kind(config, "bien_ban"),
        field_labels=FIELD_LABELS_VI,
    )


@app.route("/clients/<client_id>/edit", methods=["POST"])
@admin_required
def update_client_page(client_id: str):
    safe_client_id = normalize_client_id(client_id)
    config_path = get_client_config_path(safe_client_id)
    if not config_path.exists():
        return "Khách hàng không tồn tại.", 404

    try:
        config = json.loads(request.form.get("config_json", "{}"))
    except json.JSONDecodeError as exc:
        current = load_client_config(config_path)
        return render_template(
            "edit_client.html",
            client_id=safe_client_id,
            config=current,
            config_json=request.form.get("config_json", ""),
            template_dir=TEMPLATES_DIR / safe_client_id,
            has_phieu=has_template_kind(current, "phieu"),
            has_bien_ban=has_template_kind(current, "bien_ban"),
            field_labels=FIELD_LABELS_VI,
            error=f"JSON không hợp lệ: {exc}",
        ), 400

    if not isinstance(config.get("templates"), list):
        return "Config phải có mảng templates.", 400

    template_dir = TEMPLATES_DIR / safe_client_id
    template_dir.mkdir(parents=True, exist_ok=True)
    for index, template in enumerate(config.get("templates", [])):
        uploaded = request.files.get(f"template_file_{index}")
        if uploaded and uploaded.filename:
            if not uploaded.filename.lower().endswith((".doc", ".docx")):
                return "Template chỉ nhận file .doc hoặc .docx.", 400
            filename = saved_upload_name(uploaded, f"template_{index}_{safe_client_id}.docx")
            uploaded.save(template_dir / filename)
            template["source"] = filename

    for kind, form_name, fallback in [
        ("phieu", "new_phieu_template", f"phieu_{safe_client_id}.docx"),
        ("bien_ban", "new_bien_ban_template", f"bien_ban_{safe_client_id}.docx"),
    ]:
        uploaded = request.files.get(form_name)
        if not uploaded or not uploaded.filename:
            continue
        if has_template_kind(config, kind):
            return "Bộ form này đã có template này. Hãy dùng ô thay file template hiện tại.", 400
        if not is_word_upload(uploaded):
            return "Template chỉ nhận file .doc hoặc .docx.", 400
        filename = saved_upload_name(uploaded, fallback)
        saved_path = template_dir / filename
        uploaded.save(saved_path)
        try:
            learned_config = ai_generate_single_template_config(prepare_template(saved_path), safe_client_id, kind)
            new_template = learned_config["templates"][0]
            new_template["source"] = filename
        except Exception:
            new_template = make_template_config(safe_client_id, kind, filename)
        config["templates"].append(new_template)

    save_client_config(config_path, config)
    return redirect(url_for("edit_client_page", client_id=safe_client_id, saved="1"))


@app.route("/clients/<client_id>/delete", methods=["POST"])
@admin_required
def delete_client(client_id: str):
    safe_client_id = normalize_client_id(client_id)
    config_path = assert_inside(CLIENTS_DIR / f"{safe_client_id}.json", CLIENTS_DIR)
    template_dir = assert_inside(TEMPLATES_DIR / safe_client_id, TEMPLATES_DIR)
    output_dir = assert_inside(OUTPUT_DIR / safe_client_id, OUTPUT_DIR)

    if config_path.exists():
        config_path.unlink()
    if template_dir.exists():
        shutil.rmtree(template_dir)
    if output_dir.exists():
        shutil.rmtree(output_dir)
    return redirect(url_for("index", deleted=safe_client_id))


@app.route("/generate", methods=["GET", "POST"])
@login_required
def generate():
    if request.method == "GET":
        return redirect(url_for("index"))

    client_id = request.form.get("client_id")
    certificate = request.files.get("certificate")

    if not client_id:
        return render_template(
            "index.html",
            clients=client_ids(),
            client_list=client_details(),
            error="Vui lòng chọn Mã khách hàng cần tạo form.",
        ), 400

    if not certificate or not certificate.filename:
        return render_template(
            "index.html",
            clients=client_ids(),
            client_list=client_details(),
            error="Vui lòng chọn file Chứng thư nguồn (.docx).",
        ), 400

    if not certificate.filename.lower().endswith(".docx"):
        return render_template(
            "index.html",
            clients=client_ids(),
            client_list=client_details(),
            error="File Chứng thư nguồn chỉ nhận định dạng .docx.",
        ), 400

    config_file = CLIENTS_DIR / f"{client_id}.json"
    if not config_file.exists():
        return render_template(
            "index.html",
            clients=client_ids(),
            client_list=client_details(),
            error=f"Khách hàng '{client_id}' chưa có cấu hình mẫu. Vui lòng tạo mới mẫu trước.",
        ), 400

    cfg = load_client_config(config_file)
    run_dir = OUTPUT_DIR / client_id
    run_dir.mkdir(parents=True, exist_ok=True)

    system_settings = load_system_settings()
    use_ai = system_settings.get("enable_ai_extraction", True)
    if "use_ai" in request.form:
        use_ai = request.form.get("use_ai") == "1"
    with tempfile.TemporaryDirectory() as tmp:
        cert_path = Path(tmp) / secure_filename(certificate.filename)
        certificate.save(cert_path)
        if use_ai:
            data = extract_cert_with_ai(cert_path).as_dict()
        else:
            data = extract_certificate(cert_path).as_dict()

    generated = []
    warnings = []
    for template in cfg["templates"]:
        template_path = prepare_template(TEMPLATES_DIR / client_id / template["source"])
        output_path = run_dir / template["output"]
        missing = update_docx(template_path, output_path, template["updates"], data)
        generated.append(output_path.name)
        if missing:
            warnings.append(f"{template['source']}: không tìm thấy {', '.join(missing)}")

    return render_template(
        "result.html",
        client_id=client_id,
        data=data,
        files=generated,
        use_ai=use_ai,
        warnings=warnings,
        field_labels=FIELD_LABELS_VI,
    )


@app.get("/outputs/<client_id>/<filename>")
@login_required
def download(client_id: str, filename: str):
    return send_from_directory(OUTPUT_DIR / client_id, filename, as_attachment=True)


@app.get("/preview/<client_id>/<filename>")
@login_required
def preview(client_id: str, filename: str):
    file_path = OUTPUT_DIR / client_id / filename
    if not file_path.exists() or file_path.suffix.lower() not in [".docx", ".doc"]:
        return jsonify({"error": "File không tồn tại hoặc không đúng định dạng."}), 404
    try:
        html_content = convert_docx_to_full_html(file_path)
        return jsonify({"filename": filename, "html": html_content})
    except Exception as e:
        return jsonify({"error": f"Lỗi xem trước file: {str(e)}"}), 500


@app.errorhandler(500)
@app.errorhandler(Exception)
def handle_server_error(e):
    import traceback
    print(f"[App Exception Handled] {str(e)}")
    traceback.print_exc()
    try:
        return render_template("error.html", error=f"Đã xảy ra sự cố trên hệ thống: {str(e)}. Vui lòng thử lại."), 500
    except Exception:
        return f"<h1>Hệ thống gặp sự cố: {str(e)}</h1><p><a href='/'>Quay lại trang chủ</a></p>", 500


if __name__ == "__main__":
    from waitress import serve

    print("Serving Cert Word Sync on http://0.0.0.0:5055 via Waitress WSGI")
    serve(app, host="0.0.0.0", port=5055, threads=16, channel_timeout=120)


